"""Read immutable local snapshots into addressable units; never infer business facts."""
import csv
import hashlib
import io
import json
import re
import shutil
import subprocess
import tempfile
import xml.etree.ElementTree as ET
import zipfile
from html.parser import HTMLParser
from pathlib import Path

VERSION = 'readers-1.0'
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'


def digest(path):
    h = hashlib.sha256()
    with Path(path).open('rb') as f:
        for chunk in iter(lambda: f.read(1048576), b''):
            h.update(chunk)
    return h.hexdigest()


def unit(locator, text='', state='extracted', limitations=None, **extra):
    return dict(locator=locator, text=text, state=state,
                limitations=limitations or [], **extra)


def decode(path):
    raw = path.read_bytes()
    for encoding in ('utf-8-sig', 'utf-16' if raw[:2] in (b'\xff\xfe', b'\xfe\xff') else 'gb18030'):
        try:
            return raw.decode(encoding), encoding
        except UnicodeError:
            pass
    raise ValueError('Cannot decode text without replacement; encoding must be specified externally')


def docx(path, cache):
    rows = []
    with zipfile.ZipFile(path) as z:
        parts = sorted(n for n in z.namelist() if re.match(r'word/(document|header\d+|footer\d+|footnotes|endnotes|comments)\.xml$', n))
        for part in parts:
            tree = ET.fromstring(z.read(part))
            parents = {c: p for p in tree.iter() for c in p}

            def anchor(el):
                tags = []
                while el in parents:
                    par = parents[el]
                    siblings = [c for c in par if c.tag == el.tag]
                    tags.append(el.tag.split('}')[-1] + '[' + str(siblings.index(el) + 1) + ']')
                    el = par
                return '/'.join(reversed(tags))

            for p in tree.iter(W + 'p'):
                spans = []
                for run in p.iter(W + 'r'):
                    text = ''.join((n.text or '') if n.tag in (W + 't', W + 'delText') else '\t' if n.tag == W + 'tab' else '\n' if n.tag in (W + 'br', W + 'cr') else '' for n in run.iter())
                    if not text:
                        continue
                    flags = []
                    hide = run.find('./' + W + 'rPr/' + W + 'vanish')
                    if hide is not None and hide.get(W + 'val', 'true') not in ('0', 'false', 'off'):
                        flags.append('hidden')
                    anc = run
                    while anc in parents and anc is not p:
                        anc = parents[anc]
                        if anc.tag in (W + 'del', W + 'ins', W + 'moveFrom', W + 'moveTo'):
                            flags.append(anc.tag.split('}')[-1])
                    spans.append({'text': text, 'flags': sorted(set(flags))})
                loc = {'unit': 'block', 'part': part, 'block_path': anchor(p)}
                visible = ''.join(s['text'] for s in spans if not set(s['flags']) & {'hidden', 'del', 'moveFrom'})
                if spans:
                    rows.append(unit(loc, visible, spans=spans,
                                     limitations=['XML anchor is not a page number; layout and signatures not visually verified']))
                    for i, span in enumerate(spans, 1):
                        if span['flags']:
                            rows.append(unit(dict(loc, span=i, layer=','.join(span['flags'])), span['text'],
                                             limitations=['Non-normal text layer; do not silently treat as visible adopted text']))
        for i, name in enumerate(sorted(n for n in z.namelist() if n.startswith('word/media/') and not n.endswith('/')), 1):
            out = cache / ('media-' + str(i) + Path(name).suffix.lower())
            out.write_bytes(z.read(name))
            rows.append(unit({'unit': 'image', 'part': name}, state='manual_check',
                             limitations=['Image requires visual inspection; placement is not inferred'],
                             artifact=str(out), artifact_sha256=digest(out)))
        rels = [n for n in z.namelist() if n.endswith('.rels')]
        external = []
        for name in rels:
            for rel in ET.fromstring(z.read(name)):
                if rel.get('TargetMode') == 'External':
                    external.append({'part': name, 'target': rel.get('Target')})
        if external:
            rows.append(unit({'unit': 'block', 'part': 'relationships'}, json.dumps(external, ensure_ascii=False),
                             limitations=['External references listed only; never fetched automatically']))
    return rows


def xlsx(path, cache):
    import openpyxl
    formulas = openpyxl.load_workbook(path, data_only=False, read_only=True, keep_links=False)
    values = openpyxl.load_workbook(path, data_only=True, read_only=True, keep_links=False)
    rows = []
    try:
        for sheet in formulas:
            count = 0
            for line in sheet.iter_rows():
                for cell in line:
                    if cell.value is None:
                        continue
                    count += 1
                    formula = cell.value if cell.data_type == 'f' else None
                    value = values[sheet.title][cell.coordinate].value
                    value = value.isoformat() if hasattr(value, 'isoformat') else value
                    rows.append(unit({'unit': 'cell', 'sheet': sheet.title, 'cell_range': cell.coordinate},
                                     str(value) if value is not None else '', formula=formula,
                                     cached_or_display_value=value, recalculated_value=None, calculation_engine=None,
                                     limitations=['Cached values may be stale; reader does not recalculate formulas',
                                                  'Styles, charts, drawings, macros and external links are not visually verified']))
            rows.append(unit({'unit': 'sheet', 'sheet': sheet.title}, f'{count} populated cells extracted',
                             limitations=['Cell extraction coverage only; not business procedure coverage']))
    finally:
        formulas.close()
        values.close()
    return rows


def pdf(path, cache):
    from pypdf import PdfReader
    reader = PdfReader(path)
    if reader.is_encrypted and not reader.decrypt(''):
        raise ValueError('Encrypted PDF; password unavailable')
    rows = []
    for number, page in enumerate(reader.pages, 1):
        loc = {'unit': 'page', 'physical_page': number, 'printed_page_label': None}
        try:
            text = page.extract_text() or ''
            rows.append(unit(loc, text, 'extracted' if text.strip() else 'manual_check',
                             ['Text extraction is not page viewing; image content may remain unread']))
        except Exception as exc:
            rows.append(unit(loc, state='failed', errors=[str(exc)], limitations=['Page extraction failed']))
    return rows


def subtitles(text):
    rows = []
    timing = re.compile(r'(?m)^((?:\d{2,}:)?\d{2}:\d{2}[.,]\d{3})\s*-->\s*((?:\d{2,}:)?\d{2}:\d{2}[.,]\d{3}).*$')
    blocks = re.split(r'\n\s*\n', text.replace('\r\n', '\n').strip())
    for block in blocks:
        m = timing.search(block)
        if not m:
            if block.strip() and block.strip() != 'WEBVTT':
                rows.append(unit({'unit': 'block', 'paragraph_anchor': f'metadata-{len(rows)+1}'}, block))
            continue
        before = block[:m.start()].strip()
        content = block[m.end():].strip()
        speaker = re.match(r'\s*(?:<v\s+([^>]+)>|([^：:\n]+)[：:])', content)
        rows.append(unit({'unit': 'cue', 'cue_id': before or None,
                          'time_start': m.group(1), 'time_end': m.group(2),
                          'speaker_label': (speaker.group(1) or speaker.group(2)) if speaker else None,
                          'verified_identity': None}, content,
                         limitations=['Audio not supplied/verified; labels are not confirmed identities; preserve adjacent question and negation']))
    return rows


class PageText(HTMLParser):
    def __init__(self):
        super().__init__()
        self.skip = 0
        self.parts = []

    def handle_starttag(self, tag, attrs):
        if tag in ('script', 'style'):
            self.skip += 1
        if tag in ('p', 'div', 'li', 'br', 'h1', 'h2', 'h3', 'tr'):
            self.parts.append('\n')

    def handle_endtag(self, tag):
        if tag in ('script', 'style'):
            self.skip = max(0, self.skip - 1)

    def handle_data(self, data):
        if not self.skip:
            self.parts.append(data)


def read(path, cache):
    path, cache = Path(path), Path(cache)
    cache.mkdir(parents=True, exist_ok=True)
    provenance = []
    try:
        suffix = path.suffix.lower()
        if suffix == '.doc':
            out = cache / 'converted.docx'
            subprocess.run(['textutil', '-convert', 'docx', '-output', str(out), str(path)], check=True, capture_output=True, timeout=60)
            provenance.append({'method': 'textutil doc-to-docx', 'artifact': str(out), 'sha256': digest(out), 'parent_sha256': digest(path)})
            rows = docx(out, cache)
            for row in rows:
                row['limitations'].append('Locator belongs to converted DOCX; legacy pagination/objects may be lost')
        elif suffix == '.xls':
            with tempfile.TemporaryDirectory(prefix='evidence-xls-') as tmp:
                profile = Path(tmp) / 'profile'
                subprocess.run(['soffice', '-env:UserInstallation=' + profile.as_uri(), '--headless', '--convert-to', 'xlsx', '--outdir', tmp, str(path)], check=True, capture_output=True, timeout=60)
                converted = Path(tmp) / (path.stem + '.xlsx')
                out = cache / 'converted.xlsx'
                shutil.copyfile(converted, out)
            provenance.append({'method': 'LibreOffice xls-to-xlsx', 'artifact': str(out), 'sha256': digest(out), 'parent_sha256': digest(path)})
            rows = xlsx(out, cache)
            for row in rows:
                row['limitations'].append('Converted legacy workbook; original cached values and complex objects not preserved by guarantee')
        elif suffix == '.docx':
            rows = docx(path, cache)
        elif suffix == '.xlsx':
            rows = xlsx(path, cache)
        elif suffix == '.pdf':
            rows = pdf(path, cache)
        elif suffix in ('.png', '.jpg', '.jpeg', '.tif', '.tiff', '.webp'):
            rows = [unit({'unit': 'image', 'image': path.name}, state='manual_check', artifact=str(path), artifact_sha256=digest(path), limitations=['Visual inspection required'])]
        elif suffix in ('.txt', '.md', '.csv', '.tsv', '.srt', '.vtt', '.html', '.htm', '.json'):
            text, encoding = decode(path)
            provenance.append({'method': 'decode', 'encoding': encoding})
            if suffix in ('.srt', '.vtt'):
                rows = subtitles(text)
            elif suffix in ('.csv', '.tsv'):
                rows = [unit({'unit': 'row', 'record_number': i}, json.dumps(row, ensure_ascii=False), cells=row)
                        for i, row in enumerate(csv.reader(io.StringIO(text), delimiter='\t' if suffix == '.tsv' else ','), 1)]
            else:
                if suffix in ('.html', '.htm'):
                    parser = PageText()
                    parser.feed(text)
                    text = ''.join(parser.parts)
                rows = [unit({'unit': 'block', 'paragraph_anchor': f'line-{i}'}, line)
                        for i, line in enumerate(text.splitlines(), 1) if line.strip()]
        else:
            rows = [unit({'unit': 'file'}, state='unread', limitations=['Unsupported format; use an available format tool and register actual observations'])]
        if not rows:
            rows = [unit({'unit': 'file'}, state='manual_check', limitations=['No readable units; empty or unsupported content must be inspected'])]
        return {'reader_version': VERSION, 'units': rows, 'processing': provenance}
    except Exception as exc:
        return {'reader_version': VERSION, 'processing': provenance,
                'units': [unit({'unit': 'file'}, state='failed', errors=[f'{type(exc).__name__}: {exc}'], limitations=['No successful extraction; not evidence of absent business event'])]}
